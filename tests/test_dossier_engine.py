# -*- coding: utf-8 -*-
import os
import sys
import pytest
import docx

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from modules.dossier_engine import render_qd_du_toan_chuan_bi


def test_render_qd_du_toan_chuan_bi_defaults(tmp_path):
    out_file = str(tmp_path / "test_qd.docx")
    data = {
        "ten_cong_trinh": "Công trình Thử nghiệm An toàn 2026",
        "so_quyet_dinh": "999/QĐ-TCĐ",
        "chi_phi_lap_bcktkt_so": "50.000.000",
        "chi_phi_lap_bcktkt_chu": "Năm mươi triệu đồng"
    }

    result_path = render_qd_du_toan_chuan_bi(data, output_path=out_file)
    assert os.path.exists(result_path)

    doc = docx.Document(result_path)
    full_text = "\n".join([p.text for p in doc.paragraphs])

    assert "QUYẾT ĐỊNH" in full_text
    assert "Công trình Thử nghiệm An toàn 2026" in full_text
    assert "50.000.000 đồng" in full_text
    assert "Năm mươi triệu đồng" in full_text
    assert "Luật Xây dựng số 135/2025/QH15" in full_text
    assert "Thông tư số 101/2026/TT-BQP" in full_text


def test_render_qd_du_toan_custom_legal_bases(tmp_path):
    out_file = str(tmp_path / "test_custom_legal.docx")
    custom_bases = [
        "Căn cứ Luật Xây dựng số 135/2025/QH15;",
        "Căn cứ Nghị định số 217/2026/NĐ-CP;"
    ]
    data = {
        "ten_cong_trinh": "Công trình Nâng cấp Doanh trại S7",
        "danh_sach_can_cu": custom_bases
    }

    result_path = render_qd_du_toan_chuan_bi(data, output_path=out_file)
    assert os.path.exists(result_path)

    doc = docx.Document(result_path)
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "Căn cứ Luật Xây dựng số 135/2025/QH15;" in full_text
    assert "Căn cứ Nghị định số 217/2026/NĐ-CP;" in full_text
